"""
Lee el Excel del cierre de caja de las vendedoras ("Formato Cierre Caja <día>.xlsx",
pestañas CIERRE CAJA y CIERRE ALEGRA) para LLENAR el formulario del Cierre
diario. No guarda el archivo ni envía el cierre: la vendedora revisa y envía.
Fase 2 de docs/PLAN_EXCEDENTES_Y_CARGA_EXCEL.md (ahí está el mapa de celdas).

Las celdas se buscan por el texto de su etiqueta (ej. "Total Dinero En Caja 1",
"Transferencia Nequi") y no solo por posición, para aguantar filas movidas.
"""
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from app.exceptions import ValidationError

MAX_FILE_BYTES = 5 * 1024 * 1024

COIN_DENOMS = (50, 100, 200, 500, 1000)
BILL_DENOMS = (2000, 5000, 10000, 20000, 50000, 100000)

MONTHS_ES = {m: i for i, m in enumerate(
    ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
     'septiembre', 'octubre', 'noviembre', 'diciembre'], start=1)}

# CIERRE ALEGRA: etiqueta (columna E del formato) -> campo del formulario.
# El orden importa: "transferencia (qr) total" va antes que "transferencia (qr)".
PAYMENT_LABELS = (
    ('transferencia nequi', 'nequi_luz_helena'),
    ('transferencia daviplata', 'daviplata_jose'),
    ('transferencia (qr) total', None),  # suma de Alegra, no se usa
    ('transferencia (qr)', 'qr_julieth'),
    ('datafono addi', 'addi_datafono'),
    ('tarjeta debito', 'tarjeta_debito'),
    ('tarjeta credito', 'tarjeta_credito'),
)


def _norm(value: Any) -> str:
    """Texto en minúsculas, sin tildes y con espacios simples."""
    text = unicodedata.normalize('NFKD', str(value)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', text).strip().lower().rstrip('.:').strip()


def _amount(value: Any) -> Optional[int]:
    """Número de una celda: 3000, 3000.0, '$ 3.000' -> 3000; vacío o '-' -> None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    digits = re.sub(r'[^0-9-]', '', str(value))
    if not digits.strip('-'):
        return None
    try:
        return int(digits)
    except ValueError:
        return None


def _sheet(wb, name: str):
    for ws in wb.worksheets:
        if _norm(ws.title) == name:
            return ws
    raise ValidationError(f'El Excel no tiene la pestaña "{name.upper()}". ¿Es el formato de cierre de caja?')


def _labels(ws) -> List[Tuple[str, int, int]]:
    """(texto normalizado, fila, columna) de cada celda con texto."""
    out = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip():
                out.append((_norm(c.value), c.row, c.column))
    return out


def _find(labels, test) -> Optional[Tuple[int, int]]:
    for text, r, c in labels:
        if test(text):
            return r, c
    return None


def _value_right(ws, pos, max_offset: int) -> Optional[int]:
    """Primer valor no vacío a la derecha de la etiqueta (hasta max_offset columnas)."""
    if not pos:
        return None
    r, c = pos
    for off in range(1, max_offset + 1):
        v = ws.cell(row=r, column=c + off).value
        if v is not None and str(v).strip() != '':
            return _amount(v)
    return None


def _count_block(ws, labels, title_test, name: str) -> Dict[int, int]:
    """
    Cantidades de un cuadro de monedas/billetes: debajo del título, en la
    columna del título están las denominaciones y al lado las cantidades.
    """
    pos = _find(labels, title_test)
    if not pos:
        raise ValidationError(f'No se encontró el cuadro "{name}" en la pestaña CIERRE CAJA.')
    top, col = pos
    counts: Dict[int, int] = {}
    for r in range(top + 1, top + 30):
        label = ws.cell(row=r, column=col).value
        if isinstance(label, str) and _norm(label).startswith('total billetes y monedas'):
            break
        denom = _amount(label) if isinstance(label, (int, float)) else None
        if denom not in COIN_DENOMS + BILL_DENOMS or denom in counts:
            continue
        raw = ws.cell(row=r, column=col + 1).value
        qty = _amount(raw)
        if qty is None:
            qty = 0
        if qty < 0 or (isinstance(raw, float) and raw != int(raw)):
            raise ValidationError(f'Cantidad inválida en "{name}" para ${denom:,}: {raw}'.replace(',', '.'))
        counts[denom] = qty
    if not counts:
        raise ValidationError(f'El cuadro "{name}" no tiene denominaciones (¿se movió el formato?).')
    return counts


def _report_date(labels_alegra, ws_alegra, filename: str) -> Optional[date]:
    """Fecha del 'Reporte de ventas diarias del d/m/aaaa' o, si no, del nombre del archivo."""
    for text, r, c in labels_alegra:
        if 'reporte de ventas diarias' in text:
            m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', text)
            if m:
                try:
                    return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                except ValueError:
                    pass
            nxt = ws_alegra.cell(row=r, column=c + 1).value
            if isinstance(nxt, datetime):
                return nxt.date()
    m = re.search(r'(\d{1,2})\s+de\s+([a-z]+)\s+(?:de\s+)?(\d{4})', _norm(filename or ''))
    if m and m.group(2) in MONTHS_ES:
        try:
            return date(int(m.group(3)), MONTHS_ES[m.group(2)], int(m.group(1)))
        except ValueError:
            return None
    return None


def parse_closing_excel(content: bytes, filename: str = '', expected_date: Optional[date] = None) -> Dict[str, Any]:
    if len(content) > MAX_FILE_BYTES:
        raise ValidationError('El archivo pesa más de 5 MB: no parece el formato de cierre.')
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise ValidationError('No se pudo abrir el archivo. Debe ser el Excel del cierre guardado como .xlsx.')

    caja = _sheet(wb, 'cierre caja')
    alegra = _sheet(wb, 'cierre alegra')
    lc, la = _labels(caja), _labels(alegra)
    warnings: List[str] = []

    total = _count_block(caja, lc, lambda t: t.startswith('total dinero en caja'), 'Total Dinero En Caja 1')
    base = _count_block(caja, lc, lambda t: t.startswith('cierre de turno (base'), 'Cierre de turno (Base en Caja 1)')

    # Excedentes, gastos y préstamos (CIERRE CAJA, etiqueta en B y valor en D)
    def caja_value(test):
        return _value_right(caja, _find(lc, test), 2) or 0

    exc_datafono = caja_value(lambda t: t == 'excedente datafono')
    exc_qr = caja_value(lambda t: t.startswith('excedente qr'))
    exc_efectivo = caja_value(lambda t: t == 'excedente efectivo')
    gastos = caja_value(lambda t: t == 'gastos operativos')
    prestamos = caja_value(lambda t: t == 'prestamos')

    # Medios de pago (CIERRE ALEGRA, etiqueta en E y valor justo al lado)
    metodos = {field: 0 for _, field in PAYMENT_LABELS if field}
    found = set()
    for text, r, c in la:
        for prefix, field in PAYMENT_LABELS:
            if text.startswith(prefix):
                if field and field not in found:
                    metodos[field] = _value_right(alegra, (r, c), 1) or 0
                    found.add(field)
                break
    if not found:
        warnings.append('No se encontraron las transferencias ni las tarjetas en CIERRE ALEGRA: revísalas a mano.')

    # Notas de gastos y préstamos (CIERRE ALEGRA, columna NOTA al lado del valor)
    def note(test):
        pos = _find(la, test)
        if not pos:
            return ''
        v = alegra.cell(row=pos[0], column=pos[1] + 2).value
        return str(v).strip() if isinstance(v, str) else ''

    gastos_nota = note(lambda t: t.startswith('gastos operativos'))
    prestamos_nota = note(lambda t: t.startswith('prestamos'))

    # Lo que escribieron de Alegra: solo para avisar si no coincide
    excel_alegra = {
        'efectivo': _value_right(alegra, _find(la, lambda t: t.startswith('efectivo en alegra')), 1),
        'total': _value_right(alegra, _find(la, lambda t: t.startswith('total facturacion electronica')), 1),
    }

    excedentes = []
    if exc_efectivo:
        excedentes.append({'tipo': 'efectivo', 'subtipo': '', 'valor': exc_efectivo})
    if exc_qr:
        # El Excel no dice por cuál transferencia llegó: se pone QR
        excedentes.append({'tipo': 'qr_transferencias', 'subtipo': 'qr', 'valor': exc_qr})
        warnings.append('El Excel no dice si el excedente de transferencias llegó por QR, Nequi o Daviplata: se puso QR, cámbialo si fue otro.')
    if exc_datafono:
        excedentes.append({'tipo': 'datafono', 'subtipo': '', 'valor': exc_datafono})

    report_date = _report_date(la, alegra, filename)
    if report_date is None:
        warnings.append('No se encontró la fecha del cierre en el Excel: confirma que sea el del día.')
    elif expected_date and report_date != expected_date:
        warnings.append(f'El Excel es del {report_date.strftime("%d/%m/%Y")} y el cierre es del '
                        f'{expected_date.strftime("%d/%m/%Y")}: revisa que subiste el archivo correcto.')

    total_caja = sum(d * q for d, q in total.items())
    base_total = sum(d * q for d, q in base.items())
    for d, q in base.items():
        if q > total.get(d, 0):
            warnings.append(f'En el Excel la base deja más billetes/monedas de ${d:,} que los contados.'.replace(',', '.'))
            break

    return {
        'date': report_date.isoformat() if report_date else None,
        'coins': {str(d): total.get(d, 0) for d in COIN_DENOMS},
        'bills': {str(d): total.get(d, 0) for d in BILL_DENOMS},
        'excedentes': excedentes,
        'gastos_operativos': gastos,
        'gastos_operativos_nota': gastos_nota,
        'prestamos': prestamos,
        'prestamos_nota': prestamos_nota,
        'metodos_pago': metodos,
        'excel_alegra': excel_alegra,
        'excel_totals': {
            'total_caja': total_caja,
            'base': base_total,
            'consignar': total_caja - base_total,
        },
        'warnings': warnings,
    }
