"""
Reconstrucción de 2025 tras la anulación masiva de facturas POS
(docs/PLAN_RECONSTRUCCION_2025.md).

- Clasifica las facturas anuladas de 2025 de la copia (InvoiceFact):
  * 'masiva'  = POS anulada en la anulación masiva de oct-2026 → fue una
    VENTA REAL (cuenta como venta y sus prendas volvieron al inventario).
  * 'real'    = anulación real que se queda así: electrónica anulada, o POS
    que se volvió a facturar con las mismas prendas y el mismo total hasta
    60 minutos después, o marcada a mano (VoidOverride).
- Venta real de 2025 = vigentes + 'masiva'.
- Informe de inventario: por prenda, unidades que devolvió la anulación
  masiva y existencia antes de ella (= existencia de hoy − devueltas).
"""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Set

from sqlalchemy import func

from app.models.invoice_fact import InvoiceFact, InvoiceItemFact, InvoiceVoidItem, VoidOverride
from app.services.invoice_facts import missing_days

HISTORY_START = date(2025, 1, 1)
HISTORY_END = date(2025, 12, 31)
REPLACEMENT_WINDOW = timedelta(minutes=60)
BAG_KEYWORD = 'bolsa papel'


def _parse_dt(value: Optional[str]) -> Optional[datetime]:
    try:
        return datetime.strptime(value[:19], '%Y-%m-%d %H:%M:%S') if value else None
    except ValueError:
        return None


def coverage(store: str) -> Dict[str, Any]:
    missing = missing_days(store, HISTORY_START, HISTORY_END)
    total = (HISTORY_END - HISTORY_START).days + 1
    return {
        'start': HISTORY_START.isoformat(), 'end': HISTORY_END.isoformat(),
        'total_days': total, 'loaded_days': total - len(missing), 'missing_days': len(missing),
        'complete': not missing,
        'next_missing_day': missing[0].isoformat() if missing else None,
    }


def _signatures(store: str, start: date, end: date) -> Dict[str, tuple]:
    """Prendas (item, cantidad) de cada factura del rango, vigentes y anuladas."""
    lines = defaultdict(list)
    for model in (InvoiceItemFact, InvoiceVoidItem):
        for inv_id, item_id, qty in model.query.filter(
                model.store_code == store, model.date >= start, model.date <= end).with_entities(
                model.invoice_alegra_id, model.item_id, model.quantity):
            lines[inv_id].append((item_id or '', qty))
    return {k: tuple(sorted(v)) for k, v in lines.items()}


def classify(store: str, start: date = HISTORY_START, end: date = HISTORY_END) -> Dict[str, Any]:
    """Clasificación de cada factura anulada del rango (por defecto, 2025)."""
    facts = InvoiceFact.query.filter(
        InvoiceFact.store_code == store, InvoiceFact.date >= start, InvoiceFact.date <= end).all()
    overrides = {o.alegra_id: o for o in VoidOverride.for_store(store).all()}
    sigs = _signatures(store, start, end)
    by_day = defaultdict(list)
    for f in facts:
        by_day[f.date].append(f)

    rows = []
    for f in facts:
        if not f.voided:
            continue
        status, reason = 'masiva', 'POS anulada en la anulación masiva'
        if f.is_electronic:
            status, reason = 'real', 'Factura electrónica anulada'
        else:
            t = _parse_dt(f.issued_at)
            sig = sigs.get(f.alegra_id)
            if t and sig:
                for other in by_day[f.date]:
                    if other.alegra_id == f.alegra_id or other.total != f.total or sigs.get(other.alegra_id) != sig:
                        continue
                    ot = _parse_dt(other.issued_at)
                    if ot and t < ot <= t + REPLACEMENT_WINDOW:
                        minutes = int((ot - t).total_seconds() // 60)
                        status = 'real'
                        reason = f'Se volvió a facturar {minutes} min después (factura {other.number})'
                        break
        o = overrides.get(f.alegra_id)
        if o is not None:
            status = 'masiva' if o.counts_as_sale else 'real'
            reason = 'Marcada a mano' + (f': {o.note}' if o.note else '')
        rows.append({
            'alegra_id': f.alegra_id, 'number': f.number, 'date': f.date.isoformat(), 'issued_at': f.issued_at,
            'total': f.total, 'total_paid': f.total_paid, 'is_electronic': f.is_electronic,
            'seller_name': f.seller_name, 'status': status, 'reason': reason,
            'manual': o is not None,
        })
    return {'rows': rows}


def mass_voided_ids(store: str, start: date = HISTORY_START, end: date = HISTORY_END) -> Set[str]:
    return {r['alegra_id'] for r in classify(store, start, end)['rows'] if r['status'] == 'masiva'}


def summary(store: str) -> Dict[str, Any]:
    rows = classify(store)['rows']
    by_month = defaultdict(lambda: {'masiva_count': 0, 'masiva_total': 0, 'real_count': 0, 'real_total': 0})
    for r in rows:
        m = by_month[r['date'][:7]]
        m[f"{r['status']}_count"] += 1
        m[f"{r['status']}_total"] += r['total']
    active = InvoiceFact.query.filter(
        InvoiceFact.store_code == store, InvoiceFact.date >= HISTORY_START, InvoiceFact.date <= HISTORY_END,
        InvoiceFact.voided.is_(False))
    active_total = int(active.with_entities(func.coalesce(func.sum(InvoiceFact.total), 0)).scalar() or 0)
    masiva = [r for r in rows if r['status'] == 'masiva']
    real = [r for r in rows if r['status'] == 'real']
    return {
        'voided_count': len(rows),
        'masiva_count': len(masiva),
        'masiva_total': sum(r['total'] for r in masiva),
        'real_count': len(real),
        'real_total': sum(r['total'] for r in real),
        'active_count': active.count(),
        'active_total': active_total,
        'real_sales_total': active_total + sum(r['total'] for r in masiva),
        'by_month': [{'month': k, **v} for k, v in sorted(by_month.items())],
        'real_voids': sorted(real, key=lambda r: (r['date'], r['issued_at'] or '')),
    }


def real_sales_total(store: str, start: date, end: date) -> Optional[int]:
    """
    Venta real de un rango de 2025 desde la copia (vigentes + anulación
    masiva). None si el rango no está completo en la copia o no es de 2025.
    """
    if start < HISTORY_START or end > HISTORY_END or missing_days(store, start, end):
        return None
    facts = InvoiceFact.query.filter(
        InvoiceFact.store_code == store, InvoiceFact.date >= start, InvoiceFact.date <= end).all()
    mass = mass_voided_ids(store, start, end)
    return int(sum(f.total for f in facts if not f.voided or f.alegra_id in mass))


# ─────────────────────────────────────────────
#  Inventario
# ─────────────────────────────────────────────

def returned_units(store: str) -> Dict[str, Dict[str, Any]]:
    """Unidades por prenda que devolvió al inventario la anulación masiva."""
    mass = mass_voided_ids(store)
    out: Dict[str, Dict[str, Any]] = {}
    if not mass:
        return out
    for row in InvoiceVoidItem.query.filter(
            InvoiceVoidItem.store_code == store, InvoiceVoidItem.date >= HISTORY_START,
            InvoiceVoidItem.date <= HISTORY_END).all():
        if row.invoice_alegra_id not in mass:
            continue
        key = row.item_id or f'sin-id:{row.name}'
        r = out.setdefault(key, {'item_id': row.item_id, 'name': row.name, 'units': 0, 'invoices': set()})
        r['units'] += row.quantity
        r['invoices'].add(row.invoice_alegra_id)
    for r in out.values():
        r['invoices'] = len(r['invoices'])
    return out


def inventory_report(store: str, alegra_items: List[Dict[str, Any]]) -> Dict[str, Any]:
    """
    Cruza las unidades devueltas con la existencia y el costo de hoy en
    Alegra (`alegra_items` = AlegraClient.get_active_items()).
    """
    current = {}
    for item in alegra_items:
        inv = item.get('inventory') or {}
        if not inv:
            continue
        current[str(item.get('id'))] = {
            'name': item.get('name') or '',
            'stock': float(inv.get('availableQuantity') or 0),
            'unit_cost': float(inv.get('unitCost') or 0),
            'reference': (item.get('reference') or {}).get('reference') if isinstance(item.get('reference'), dict) else item.get('reference'),
        }
    value_now = sum(c['stock'] * c['unit_cost'] for c in current.values())

    rows = []
    for key, r in returned_units(store).items():
        c = current.get(str(r['item_id'])) if r['item_id'] else None
        stock = c['stock'] if c else None
        cost = c['unit_cost'] if c else 0
        before = (stock - r['units']) if stock is not None else None
        flag = None
        if c is None:
            flag = 'No está activa en Alegra (o no maneja inventario)'
        elif before < 0:
            flag = 'Da negativo: la existencia de hoy es menor que lo devuelto; revisar en físico'
        units_to_remove = min(r['units'], stock) if stock is not None and stock > 0 else 0
        rows.append({
            'item_id': r['item_id'], 'name': c['name'] if c else r['name'], 'reference': c['reference'] if c else None,
            'units_returned': r['units'], 'invoices': r['invoices'],
            'stock_now': stock, 'stock_before': max(before, 0) if before is not None else None,
            'units_to_remove': units_to_remove,
            'unit_cost': cost, 'value_to_remove': units_to_remove * cost,
            'is_bag': BAG_KEYWORD in (r['name'] or '').lower(),
            'flag': flag,
        })
    rows.sort(key=lambda x: -x['value_to_remove'])
    value_to_remove = sum(r['value_to_remove'] for r in rows)
    return {
        'rows': rows,
        'items_count': len(rows),
        'units_returned': sum(r['units_returned'] for r in rows),
        'units_to_remove': sum(r['units_to_remove'] for r in rows),
        'value_now': value_now,
        'value_to_remove': value_to_remove,
        'value_before': value_now - value_to_remove,
        'flagged': sum(1 for r in rows if r['flag']),
    }


def inventory_excel(report: Dict[str, Any], store_name: str) -> bytes:
    """Excel del informe para revisarlo con el contador."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = 'Ajuste propuesto'
    ws.append([f'{store_name} - Inventario antes de la anulación masiva de facturas POS 2025'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([f'Generado: {datetime.now().strftime("%Y-%m-%d %H:%M")}. Existencia antes = existencia hoy en Alegra − unidades que devolvió la anulación masiva.'])
    ws.append(['Valor del inventario hoy (Alegra, a costo)', round(report['value_now'])])
    ws.append(['Valor a retirar (ajuste de salida)', round(report['value_to_remove'])])
    ws.append(['Valor del inventario antes de la anulación (estimado)', round(report['value_before'])])
    ws.append([])
    headers = ['ID Alegra', 'Referencia', 'Prenda', 'Unidades devueltas por la anulación', 'Facturas',
               'Existencia hoy en Alegra', 'Existencia antes de la anulación', 'Unidades a retirar (ajuste)',
               'Costo unitario', 'Valor a retirar', 'Revisar']
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F2937')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for r in report['rows']:
        ws.append([r['item_id'], r['reference'], r['name'], r['units_returned'], r['invoices'],
                   r['stock_now'], r['stock_before'], r['units_to_remove'], round(r['unit_cost']),
                   round(r['value_to_remove']), r['flag'] or ''])
    widths = [10, 16, 44, 16, 10, 14, 16, 16, 14, 16, 48]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
