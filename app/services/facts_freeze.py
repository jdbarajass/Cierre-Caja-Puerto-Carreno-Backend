"""
Blindaje de la copia de facturas (docs/PLAN_BLINDAJE_COPIA.md).

Para cuando se anulen en Alegra, de forma masiva, facturas de un año ya
guardado (como pasó con las POS de 2025): el usuario CONGELA la copia hasta
una fecha justo antes de esa anulación. Desde ese momento:

- Los días congelados nunca se vuelven a descargar de Alegra (ni con la carga
  nocturna, ni al subir FACT_VERSION): la copia queda como "la verdad".
- En las pantallas que leen Alegra en vivo, las facturas de días congelados
  que Alegra muestra anuladas pero que en la copia estaban vigentes vuelven
  a contar como venta (`frozen_valid_ids`, usado por history_2025.revive_*).
- Las anulaciones REALES hechas antes de congelar ya están en la copia: la
  carga nocturna repasa los últimos 3 días y, antes de congelar, el "repaso
  completo" vuelve a descargar todo el rango una vez.

Sin congelar no cambia nada (todas las funciones devuelven vacío / None).
"""
import json
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Set, Tuple

from app.models.invoice_fact import InvoiceFact, InvoiceItemFact, InvoiceSyncDay
from app.models.user import db

SETTING = 'invoice_facts_freeze'
FACTS_START = date(2026, 1, 1)   # primer día de la copia (app/routes/invoice_facts.py)


class FreezeError(Exception):
    pass


# ─── Estado ──────────────────────────────────────────────────────────────────

def _row(store: str):
    from app.models.app_setting import AppSetting
    from app.stores import store_setting_key
    return AppSetting.query.get(store_setting_key(SETTING, store))


def state(store: str) -> Optional[Dict[str, Any]]:
    row = _row(store)
    return json.loads(row.value) if row and row.value else None


def frozen_until(store: str) -> Optional[date]:
    s = state(store)
    return date.fromisoformat(s['until']) if s and s.get('until') else None


def frozen_part(store: str, start: date, end: date) -> Optional[Tuple[date, date]]:
    """Parte del rango que está congelada (None si nada)."""
    until = frozen_until(store)
    if until is None:
        return None
    s, e = max(start, FACTS_START), min(end, until)
    return (s, e) if s <= e else None


def is_frozen_day(store: str, day: date, until: Optional[date] = None) -> bool:
    until = until if until is not None else frozen_until(store)
    return until is not None and FACTS_START <= day <= until


def unfrozen_parts(store: str, start: date, end: date) -> List[Tuple[date, date]]:
    """Partes del rango que NO están congeladas (lo que se sigue pidiendo a Alegra)."""
    part = frozen_part(store, start, end)
    if part is None:
        return [(start, end)] if start <= end else []
    out = []
    if start < part[0]:
        out.append((start, date.fromordinal(part[0].toordinal() - 1)))
    if part[1] < end:
        out.append((date.fromordinal(part[1].toordinal() + 1), end))
    return out


# ─── Lo que usan las pantallas ───────────────────────────────────────────────

def frozen_valid_ids(store: str, ids) -> Set[str]:
    """De esos ids, los que en la copia congelada estaban vigentes."""
    until = frozen_until(store)
    ids = [str(i) for i in ids]
    if until is None or not ids:
        return set()
    rows = InvoiceFact.query.filter(
        InvoiceFact.store_code == store, InvoiceFact.alegra_id.in_(ids), InvoiceFact.voided.is_(False),
        InvoiceFact.date >= FACTS_START, InvoiceFact.date <= until).with_entities(InvoiceFact.alegra_id)
    return {r.alegra_id for r in rows}


def frozen_valid_ids_in_range(store: str, start: date, end: date) -> Set[str]:
    """Todas las vigentes de la parte congelada del rango (para el comparativo, que corre en hilos)."""
    part = frozen_part(store, start, end)
    if part is None:
        return set()
    rows = InvoiceFact.query.filter(
        InvoiceFact.store_code == store, InvoiceFact.voided.is_(False),
        InvoiceFact.date >= part[0], InvoiceFact.date <= part[1]).with_entities(InvoiceFact.alegra_id)
    return {r.alegra_id for r in rows}


def copy_sales_total(store: str, start: date, end: date) -> Optional[int]:
    """Venta del rango según la copia, si TODO el rango está congelado (si no, None)."""
    part = frozen_part(store, start, end)
    if part != (start, end):
        return None
    total = db.session.query(db.func.coalesce(db.func.sum(InvoiceFact.total), 0)).filter(
        InvoiceFact.store_code == store, InvoiceFact.voided.is_(False),
        InvoiceFact.date >= start, InvoiceFact.date <= end).scalar()
    return int(total or 0)


def copy_client_rows(store: str, start: date, end: date) -> List[Dict[str, Any]]:
    """Compras por cliente de la parte congelada del rango, con el formato de sales-by-client de Alegra."""
    part = frozen_part(store, start, end)
    if part is None:
        return []
    out: Dict[str, Dict[str, Any]] = {}
    for f in InvoiceFact.query.filter(
            InvoiceFact.store_code == store, InvoiceFact.voided.is_(False),
            InvoiceFact.date >= part[0], InvoiceFact.date <= part[1]):
        if not f.client_id:
            continue
        r = out.setdefault(f.client_id, {'idLocal': f.client_id, 'clientName': f.client_name,
                                         'identification': f.client_identification,
                                         'totalDocuments': 0, 'subtotal': 0, 'discount': 0, 'total': 0})
        r['totalDocuments'] += 1
        r['subtotal'] += f.subtotal or 0
        r['discount'] += f.discount or 0
        r['total'] += f.total or 0
    return list(out.values())


# ─── Repaso y congelamiento ──────────────────────────────────────────────────

def status(store: str, until: date) -> Dict[str, Any]:
    """Qué tan lista está la copia para congelarla hasta `until`."""
    from app.services.invoice_facts import missing_days, missing_item_days, outdated_days
    current = frozen_until(store)
    # Lo ya congelado no se vuelve a descargar: no cuenta como pendiente
    start = FACTS_START if current is None else max(FACTS_START, date.fromordinal(current.toordinal() + 1))
    pending = sorted(set(missing_days(store, start, until)) | set(missing_item_days(store, start, until))
                     | set(outdated_days(store, start, until))) if start <= until else []
    facts = InvoiceFact.query.filter(InvoiceFact.store_code == store, InvoiceFact.date >= FACTS_START,
                                     InvoiceFact.date <= until)
    s = state(store) or {}
    return {
        'until': until.isoformat(),
        'total_days': (until - FACTS_START).days + 1,
        'pending_days': len(pending),
        'next_pending': [d.isoformat() for d in pending[-5:]],
        'invoices': facts.count(),
        'active_invoices': facts.filter(InvoiceFact.voided.is_(False)).count(),
        'voided_invoices': facts.filter(InvoiceFact.voided.is_(True)).count(),
        'active_total': int(facts.filter(InvoiceFact.voided.is_(False)).with_entities(
            db.func.coalesce(db.func.sum(InvoiceFact.total), 0)).scalar() or 0),
        'items': InvoiceItemFact.query.filter(InvoiceItemFact.store_code == store, InvoiceItemFact.date >= FACTS_START,
                                              InvoiceItemFact.date <= until).count(),
        'review': s.get('review'),
        'frozen': s if s.get('until') else None,
        'ready': not pending,
    }


def start_review(store: str, until: date, user_id: Optional[int]) -> int:
    """
    Marca los días ya cargados (no congelados) hasta `until` para volver a
    descargarlos una vez (la carga nocturna y el botón de tandas los toman):
    así la copia recoge las anulaciones reales hechas después de los 3 días
    que repasa cada noche. Devuelve cuántos días quedaron por repasar.
    """
    current = frozen_until(store)
    start = FACTS_START if current is None else date.fromordinal(current.toordinal() + 1)
    if start > until:
        return 0
    n = InvoiceSyncDay.query.filter(
        InvoiceSyncDay.store_code == store, InvoiceSyncDay.date >= start, InvoiceSyncDay.date <= until
    ).update({InvoiceSyncDay.fact_version: 0}, synchronize_session=False)
    _save(store, {**(state(store) or {}), 'review': {
        'until': until.isoformat(), 'started_at': datetime.utcnow().isoformat() + 'Z', 'days': n, 'user_id': user_id}})
    return n


def freeze(store: str, until: date, today: date, user_id: Optional[int]) -> Dict[str, Any]:
    if until < FACTS_START:
        raise FreezeError(f'La copia empieza el {FACTS_START.isoformat()}')
    if until >= today:
        raise FreezeError('Solo se pueden congelar días que ya pasaron (hasta ayer)')
    current = frozen_until(store)
    if current is not None and until <= current:
        raise FreezeError(f'Ya está congelada hasta el {current.isoformat()}: solo se puede extender')
    st = status(store, until)
    if not st['ready']:
        raise FreezeError(f'Faltan {st["pending_days"]} día(s) por cargar o repasar antes de congelar')
    s = state(store) or {}
    history = s.get('history', [])
    if s.get('until'):
        history.append({k: s[k] for k in ('until', 'frozen_at', 'user_id', 'invoices', 'active_total') if k in s})
    new = {'until': until.isoformat(), 'frozen_at': datetime.utcnow().isoformat() + 'Z', 'user_id': user_id,
           'invoices': st['invoices'], 'active_total': st['active_total'], 'history': history}
    _save(store, new)
    return new


def _save(store: str, value: Dict[str, Any]) -> None:
    from app.models.app_setting import AppSetting
    from app.stores import store_setting_key
    row = _row(store)
    if row is None:
        row = AppSetting(key=store_setting_key(SETTING, store))
        db.session.add(row)
    row.value = json.dumps(value, ensure_ascii=False)
    row.updated_at = datetime.utcnow()
    db.session.commit()


# ─── Respaldo en Excel ───────────────────────────────────────────────────────

def backup_excel(store: str, store_name: str, start: date, end: date) -> bytes:
    """Facturas y prendas guardadas del rango (copia fuera de la plataforma)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from app.models.invoice_fact import InvoiceVoidItem

    def header(ws, cols):
        ws.append(cols)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F2937')
        ws.freeze_panes = 'A2'

    wb = Workbook()
    ws = wb.active
    ws.title = 'Facturas'
    header(ws, ['ID Alegra', 'Fecha', 'Hora', 'Número', 'Electrónica', 'Anulada', 'Cliente', 'Cédula',
                'ID cliente', 'Vendedora', 'ID vendedora', 'Subtotal', 'Descuento', 'Total', 'Pagado'])
    for f in InvoiceFact.query.filter(InvoiceFact.store_code == store, InvoiceFact.date >= start,
                                      InvoiceFact.date <= end).order_by(InvoiceFact.date, InvoiceFact.issued_at):
        ws.append([f.alegra_id, f.date, (f.issued_at or '')[11:], f.number,
                   'Sí' if f.is_electronic else ('No' if f.is_electronic is False else ''),
                   'Sí' if f.voided else 'No', f.client_name, f.client_identification, f.client_id,
                   f.seller_name, f.seller_id, f.subtotal, f.discount, f.total, f.total_paid])
    for col, w in zip('ABCDEFGHIJKLMNO', (10, 11, 9, 12, 11, 9, 34, 15, 10, 24, 11, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet('Prendas vendidas')
    header(ws, ['ID factura', 'Fecha', 'ID prenda', 'Prenda', 'Cantidad', 'Precio', '% descuento', 'Total', 'Vendedora'])
    for r in InvoiceItemFact.query.filter(InvoiceItemFact.store_code == store, InvoiceItemFact.date >= start,
                                          InvoiceItemFact.date <= end).order_by(InvoiceItemFact.date):
        ws.append([r.invoice_alegra_id, r.date, r.item_id, r.name, r.quantity, r.unit_price, r.discount_pct,
                   r.total, r.seller_name])
    for col, w in zip('ABCDEFGHI', (11, 11, 10, 46, 9, 12, 11, 12, 24)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet('Prendas de anuladas')
    header(ws, ['ID factura', 'Fecha', 'ID prenda', 'Prenda', 'Cantidad', 'Precio', 'Total'])
    for r in InvoiceVoidItem.query.filter(InvoiceVoidItem.store_code == store, InvoiceVoidItem.date >= start,
                                          InvoiceVoidItem.date <= end).order_by(InvoiceVoidItem.date):
        ws.append([r.invoice_alegra_id, r.date, r.item_id, r.name, r.quantity, r.unit_price, r.total])
    for col, w in zip('ABCDEFG', (11, 11, 10, 46, 9, 12, 12)):
        ws.column_dimensions[col].width = w

    info = wb.create_sheet('Información', 0)
    info.append([f'{store_name} - Respaldo de la copia de facturas'])
    info['A1'].font = Font(bold=True, size=13)
    info.append(['Rango', f'{start.isoformat()} a {end.isoformat()}'])
    info.append(['Generado', datetime.now().strftime('%Y-%m-%d %H:%M')])
    until = frozen_until(store)
    info.append(['Copia congelada hasta', until.isoformat() if until else 'No congelada'])
    info.append(['Nota', 'Copia de la plataforma (no de Alegra). "Anulada" = como estaba en la copia.'])
    info.column_dimensions['A'].width = 24
    info.column_dimensions['B'].width = 70

    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
