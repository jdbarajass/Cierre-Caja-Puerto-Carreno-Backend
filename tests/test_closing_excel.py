"""
Carga del Excel del cierre para llenar el formulario (Fase 2 de
docs/PLAN_EXCEDENTES_Y_CARGA_EXCEL.md). El Excel se arma aquí con openpyxl
imitando el formato de las vendedoras (el real no se versiona).
"""
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.exceptions import ValidationError
from app.services.closing_excel import parse_closing_excel

COUNTS = {100: 0, 200: 15, 500: 0, 1000: 0, 2000: 18, 5000: 3, 10000: 17, 20000: 8, 50000: 4, 100000: 6}
BASE = {200: 15, 2000: 16, 5000: 3, 10000: 16, 20000: 7, 50000: 2}


def _workbook(shift=0, qr_excedente=None, report='Reporte de ventas diarias del  7/10/2026'):
    """Formato del 7-oct-2026 (con `shift` filas extra arriba en CIERRE CAJA)."""
    wb = Workbook()
    wb.active.title = 'Tarjetas Regalo'
    caja = wb.create_sheet('CIERRE CAJA')
    s = shift
    caja.cell(row=2 + s, column=2, value='Total Dinero En Caja 1')
    caja.cell(row=2 + s, column=6, value='Cierre de turno (Base en Caja 1 = $450,000)')
    rows = {100: 5, 200: 6, 500: 7, 1000: 8, 2000: 13, 5000: 14, 10000: 15, 20000: 16, 50000: 17, 100000: 18}
    for denom, r in rows.items():
        caja.cell(row=r + s, column=2, value=denom)
        caja.cell(row=r + s, column=3, value=COUNTS[denom])
        caja.cell(row=r + s, column=6, value=denom)
        caja.cell(row=r + s, column=7, value=BASE.get(denom, 0))
    caja.cell(row=4 + s, column=2, value='Denominación')
    caja.cell(row=20 + s, column=2, value='TOTAL BILLETES Y MONEDAS')
    caja.cell(row=20 + s, column=6, value='TOTAL BILLETES Y MONEDAS')
    for r, label, value in ((23, 'Excedente Datafono', 0), (24, 'Excedente QR/Transferencias', qr_excedente),
                            (25, 'Excedente Efectivo', 600), (26, 'Gastos Operativos', 3000), (27, 'Préstamos', 0)):
        caja.cell(row=r + s, column=2, value=label)
        if value is not None:
            caja.cell(row=r + s, column=4, value=value)
    caja.cell(row=32 + s, column=2, value='Valor Efectivo Sumando Gastos Operativos')
    caja.cell(row=32 + s, column=4, value=737000)

    alegra = wb.create_sheet('CIERRE ALEGRA')
    alegra.cell(row=3, column=2, value=report)
    alegra.cell(row=13, column=2, value='Total Facturación Electrónica')
    alegra.cell(row=13, column=3, value=1605650)
    for r, label, value in ((10, 'Efectivo en Alegra(Jhonatan)', 736400), (11, 'Transferencia Nequi (Luz Helena)', 0),
                            (12, 'Transferencia Daviplata (Jose Barajas)', 0), (13, 'Transferencia (QR) (Jose Barajas)', 599550),
                            (14, 'Transferencia (QR) Total (Alegra)', 599550), (15, 'Datafono Addi (Cristhian)', None),
                            (16, 'Tarjeta débito en Alegra(Cristhian)', 269700), (17, 'Tarjeta Crédito en Alegra(Cristhian)', 0)):
        alegra.cell(row=r, column=5, value=label)
        if value is not None:
            alegra.cell(row=r, column=6, value=value)
    alegra.cell(row=12, column=7, value=600)          # otro cuadro al lado: no se debe leer
    alegra.cell(row=15, column=7, value='Cae en Data pero en Alegra es Trx')
    alegra.cell(row=28, column=2, value='Gastos operativos ()')
    alegra.cell(row=28, column=4, value='bolsas ')
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_lee_el_formato_del_7_de_octubre():
    r = parse_closing_excel(_workbook(), 'Formato Cierre Caja 07 de OCTUBRE 2026.xlsx', date(2026, 10, 7))
    assert r['date'] == '2026-10-07' and r['warnings'] == []
    assert r['coins'] == {'50': 0, '100': 0, '200': 15, '500': 0, '1000': 0}
    assert r['bills'] == {'2000': 18, '5000': 3, '10000': 17, '20000': 8, '50000': 4, '100000': 6}
    assert r['excel_totals'] == {'total_caja': 1184000, 'base': 450000, 'consignar': 734000}
    assert r['excedentes'] == [{'tipo': 'efectivo', 'subtipo': '', 'valor': 600}]
    assert (r['gastos_operativos'], r['gastos_operativos_nota'], r['prestamos']) == (3000, 'bolsas', 0)
    assert r['metodos_pago'] == {'nequi_luz_helena': 0, 'daviplata_jose': 0, 'qr_julieth': 599550,
                                 'addi_datafono': 0, 'tarjeta_debito': 269700, 'tarjeta_credito': 0}
    assert r['excel_alegra'] == {'efectivo': 736400, 'total': 1605650}


def test_filas_movidas_excedente_qr_y_otra_fecha():
    r = parse_closing_excel(_workbook(shift=3, qr_excedente=5000), 'cierre.xlsx', date(2026, 10, 8))
    assert r['bills']['100000'] == 6 and r['excel_totals']['consignar'] == 734000
    assert {'tipo': 'qr_transferencias', 'subtipo': 'qr', 'valor': 5000} in r['excedentes']
    assert any('QR, Nequi o Daviplata' in w for w in r['warnings'])
    assert any('07/10/2026' in w and '08/10/2026' in w for w in r['warnings'])


def test_fecha_desde_el_nombre_del_archivo():
    r = parse_closing_excel(_workbook(report='Reporte de ventas diarias'), 'Formato Cierre Caja 07 de OCTUBRE 2026.xlsx')
    assert r['date'] == '2026-10-07'


def test_archivo_que_no_es_el_formato():
    with pytest.raises(ValidationError):
        parse_closing_excel(b'esto no es un excel', 'x.xlsx')
    wb = Workbook()
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ValidationError) as exc:
        parse_closing_excel(buf.getvalue(), 'x.xlsx')
    assert 'CIERRE CAJA' in exc.value.message


# ─── Endpoint ────────────────────────────────────────────────────────────────

@pytest.fixture
def app(tmp_path):
    from app.config import TestingConfig

    class Cfg(TestingConfig):
        SQLALCHEMY_DATABASE_URI = f"sqlite:///{(tmp_path / 'excel.db').as_posix()}"
        SQLALCHEMY_ENGINE_OPTIONS = {}

    from app import create_app
    return create_app(Cfg)


def _headers(app, role):
    from app.services.jwt_service import JWTService
    with app.app_context():
        return {'Authorization': f'Bearer {JWTService.generate_token(2, "v@test.com", role, None)}'}


def test_endpoint_llena_sin_guardar_nada(app, client):
    from app.models.cash_closing import CashClosing
    h = _headers(app, 'sales')
    resp = client.post('/api/cash_closing/parse-excel', headers=h, content_type='multipart/form-data',
                       data={'file': (BytesIO(_workbook()), 'Formato Cierre Caja 07 de OCTUBRE 2026.xlsx'),
                             'date': '2026-10-07'})
    body = resp.get_json()
    assert resp.status_code == 200 and body['success'] and body['metodos_pago']['qr_julieth'] == 599550
    with app.app_context():
        assert CashClosing.query.count() == 0   # solo lee: no crea cierre

    bad = client.post('/api/cash_closing/parse-excel', headers=h, content_type='multipart/form-data',
                      data={'file': (BytesIO(b'a,b'), 'cierre.csv')})
    assert bad.status_code == 400 and '.xlsx' in bad.get_json()['error']
    assert client.post('/api/cash_closing/parse-excel', headers=h, data={}).status_code == 400
    assert client.post('/api/cash_closing/parse-excel', content_type='multipart/form-data',
                       data={'file': (BytesIO(_workbook()), 'c.xlsx')}).status_code == 401
