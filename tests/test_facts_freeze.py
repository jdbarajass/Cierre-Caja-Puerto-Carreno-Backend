"""
Blindaje de la copia de facturas (docs/PLAN_BLINDAJE_COPIA.md): se congela
la copia y luego Alegra anula facturas de forma masiva. La copia no cambia y
las pantallas cuentan la venta real; las anulaciones reales de antes de
congelar siguen anuladas. Sin congelar, nada cambia.

Facturas con el formato real de /api/v1/invoices. Base SQLite temporal, sin red.
"""
from datetime import date, datetime

import pytest

from app.config import TestingConfig

START = date(2026, 3, 1)   # la copia de la prueba empieza aquí (FACTS_START real: 1-ene-2026)


@pytest.fixture
def app(tmp_path, monkeypatch):
    class FreezeTestConfig(TestingConfig):
        SQLALCHEMY_DATABASE_URI = f"sqlite:///{(tmp_path / 'freeze.db').as_posix()}"
        SQLALCHEMY_ENGINE_OPTIONS = {}

    import app.services.facts_freeze as freeze_svc
    monkeypatch.setattr(freeze_svc, 'FACTS_START', START)
    from app import create_app
    return create_app(FreezeTestConfig)


@pytest.fixture
def h(app):
    from app.services.jwt_service import JWTService
    with app.app_context():
        token = JWTService.generate_token(1, 'admin1@test.com', 'admin', None)
    return lambda store=None: {'Authorization': f'Bearer {token}', **({'X-Store': store} if store else {})}


@pytest.fixture(autouse=True)
def today(monkeypatch):
    import app.routes.facts_freeze as routes
    monkeypatch.setattr(routes, 'get_colombia_now', lambda: datetime(2026, 3, 10, 12, 0))


def inv(id_, day, total, status='closed', method='CASH'):
    return {
        'id': str(id_), 'date': day, 'datetime': f'{day} 10:00:00', 'status': status, 'paymentMethod': method,
        'subtotal': total, 'discount': 0, 'tax': 0, 'total': total,
        'totalPaid': total if status != 'void' else 0, 'balance': 0,
        'numberTemplate': {'id': '16', 'prefix': 'KPC', 'number': str(id_), 'fullNumber': f'KPC{id_}', 'isElectronic': True},
        'client': {'id': '50', 'name': 'ANA', 'identification': '1001'},
        'seller': {'id': '7', 'name': 'MONICA VARGAS'},
        'items': [{'id': '1183', 'name': 'JEAN HOMBRE 109900 / 105110990034', 'price': total, 'discount': 0,
                   'quantity': 1, 'total': total}],
    }


# 2-mar: dos ventas y una anulación real (se anuló el mismo día)
ANTES = {
    '2026-03-01': [inv(10, '2026-03-01', 50000)],
    '2026-03-02': [inv(11, '2026-03-02', 100000), inv(12, '2026-03-02', 200000), inv(13, '2026-03-02', 70000, status='void')],
}
# Después: la anulación masiva anula todo en Alegra
DESPUES = {d: [dict(i, status='void', totalPaid=0) for i in rows] for d, rows in ANTES.items()}


class FakeAlegra:
    def __init__(self, by_day):
        self.by_day = by_day

    def get_invoices_by_date(self, day):
        return self.by_day.get(day, [])

    def get_all_invoices_in_range(self, start, end):
        return [i for d, rows in sorted(self.by_day.items()) if start <= d <= end for i in rows]


def _load(app, by_day=ANTES):
    from app.services.invoice_facts import sync_day
    with app.app_context():
        for d in by_day:
            sync_day(FakeAlegra(by_day), 'carreno', date.fromisoformat(d))


def _freeze(client, h, until='2026-03-02', confirm='CONGELAR'):
    return client.post('/api/facts-freeze/freeze', headers=h(), json={'until': until, 'confirm': confirm})


def _active_total(app):
    from app.models.invoice_fact import InvoiceFact
    with app.app_context():
        return sum(f.total for f in InvoiceFact.query.filter_by(store_code='carreno', voided=False))


def test_sin_congelar_todo_sigue_igual(app):
    """Sin congelar, una recarga refleja lo que diga Alegra (comportamiento de siempre)."""
    from app.services.history_2025 import real_sales_total, revive_mass_voided
    _load(app)
    _load(app, DESPUES)
    assert _active_total(app) == 0
    with app.app_context():
        assert real_sales_total('carreno', date(2026, 3, 1), date(2026, 3, 2)) is None
        assert revive_mass_voided(DESPUES['2026-03-02'], 'carreno') == DESPUES['2026-03-02']


def test_congelar_pide_confirmacion_dias_completos_y_solo_extender(app, client, h):
    _load(app, {'2026-03-01': ANTES['2026-03-01']})          # falta el 2-mar
    assert 'CONGELAR' in _freeze(client, h, confirm='si').get_json()['message']
    resp = _freeze(client, h)
    assert resp.status_code == 400 and 'Faltan 1 día' in resp.get_json()['message']
    _load(app)
    assert _freeze(client, h, until='2026-03-10').status_code == 400        # hoy: no
    st = client.get('/api/facts-freeze/status?until=2026-03-02', headers=h()).get_json()
    assert st['ready'] is True and st['active_total'] == 350000 and st['voided_invoices'] == 1
    assert _freeze(client, h).status_code == 200
    assert _freeze(client, h, until='2026-03-01').status_code == 400         # no se reduce
    assert client.get('/api/facts-freeze/status?until=2026-03-02', headers=h()).get_json()['frozen']['until'] == '2026-03-02'
    # la otra tienda no queda congelada
    assert client.get('/api/facts-freeze/status?until=2026-03-02', headers=h('primavera')).get_json()['frozen'] is None


def test_anulacion_masiva_despues_de_congelar(app, client, h, monkeypatch):
    import app.services.invoice_facts as facts_svc
    from app.services.history_2025 import live_revive_ids, real_sales_total, revive_mass_voided
    from app.utils.formatters import filter_voided_invoices
    _load(app)
    assert _freeze(client, h).status_code == 200

    # La anulación masiva llega en una recarga (y aunque suba FACT_VERSION): la copia no cambia
    _load(app, DESPUES)
    monkeypatch.setattr(facts_svc, 'FACT_VERSION', 99)
    with app.app_context():
        assert facts_svc.pending_days('carreno', START, date(2026, 3, 2)) == []
    assert _active_total(app) == 350000

    with app.app_context():
        # Metas / Dashboard / año anterior: venta real desde la copia
        assert real_sales_total('carreno', date(2026, 3, 1), date(2026, 3, 2)) == 350000
        assert real_sales_total('carreno', date(2026, 3, 1), date(2026, 3, 9)) is None   # el 3-mar en adelante no está
        # Pantallas en vivo: reviven las vigentes; la anulación real (13) sigue anulada
        revived = revive_mass_voided(DESPUES['2026-03-02'], 'carreno')
        active = filter_voided_invoices(revived)['active_invoices']
        assert sorted(i['id'] for i in active) == ['11', '12'] and all(i['mass_voided'] for i in active)
        assert live_revive_ids('carreno', date(2026, 3, 1), date(2026, 3, 31)) == {'10', '11', '12'}

    # Documentos de Venta (Totales usa el mismo endpoint)
    import app.routes.direct_api as direct_routes

    class Direct:
        def get_all_invoices_for_date_range(self, from_date, to_date):
            return {'success': True, 'data': DESPUES['2026-03-02'], 'metadata': {}}
    monkeypatch.setattr(direct_routes, 'get_alegra_direct_client', lambda *a, **k: Direct())
    data = client.get('/api/direct/sales/documents?from=2026-03-02&to=2026-03-02', headers=h()).get_json()
    assert sum(i['total'] for i in data['data']) == 300000 and data['voided']['count'] == 1

    # Comparativo de tiendas (corre en hilos: ids calculados antes)
    from app.routes.stores import _sales_metrics
    from app.services.alegra_client import AlegraClient

    class Alegra(FakeAlegra):
        build_sales_summary = AlegraClient.build_sales_summary
        last_failed_days = []
        username = 'test'
    with app.app_context():
        ids = live_revive_ids('carreno', date(2026, 3, 1), date(2026, 3, 2))
    sales = _sales_metrics(Alegra(DESPUES), date(2026, 3, 1), date(2026, 3, 2), ids)
    assert sales['total'] == 350000 and sales['invoices'] == 3


def test_pagos_de_dias_congelados_no_se_reemplazan(app, client, h):
    from app.models.month_sheet import PaymentFact
    from app.models.user import db
    from app.services.payment_facts import sync_payments
    _load(app)
    with app.app_context():
        db.session.add(PaymentFact(store_code='carreno', payment_id='p1', invoice_id='11', invoice_date=date(2026, 3, 2),
                                   payment_date=date(2026, 3, 2), medio='efectivo', amount=100000))
        db.session.commit()
    assert _freeze(client, h).status_code == 200

    class NoPayments:      # los recibos se anularon en la anulación masiva
        def get_payments_page(self, start, limit=30):
            return []
    with app.app_context():
        sync_payments(NoPayments(), 'carreno', date(2026, 3, 1))
        assert PaymentFact.query.filter_by(store_code='carreno').count() == 1


def test_repaso_completo_recoge_anulaciones_reales_tardias(app, client, h):
    """Antes de congelar: una anulación real hecha días después se recoge con el repaso."""
    from app.services.invoice_facts import sync_range
    _load(app)
    tardia = {**ANTES, '2026-03-01': [inv(10, '2026-03-01', 50000, status='void')]}
    resp = client.post('/api/facts-freeze/review', headers=h(), json={'until': '2026-03-02'})
    assert resp.status_code == 200 and resp.get_json()['days'] == 2 and resp.get_json()['ready'] is False
    with app.app_context():
        from app.services.invoice_facts import pending_days
        sync_range(FakeAlegra(tardia), 'carreno', pending_days('carreno', START, date(2026, 3, 2)))
    assert _active_total(app) == 300000
    assert client.get('/api/facts-freeze/status?until=2026-03-02', headers=h()).get_json()['ready'] is True


def test_clientes_partes_y_respaldo(app, client, h):
    from app.services.facts_freeze import copy_client_rows, unfrozen_parts
    _load(app)
    with app.app_context():
        assert unfrozen_parts('carreno', date(2025, 12, 1), date(2026, 3, 31)) == [(date(2025, 12, 1), date(2026, 3, 31))]
    assert _freeze(client, h).status_code == 200
    with app.app_context():
        assert unfrozen_parts('carreno', date(2025, 12, 1), date(2026, 3, 31)) == [
            (date(2025, 12, 1), date(2026, 2, 28)), (date(2026, 3, 3), date(2026, 3, 31))]
        rows = copy_client_rows('carreno', date(2026, 1, 1), date(2026, 3, 31))
    assert rows == [{'idLocal': '50', 'clientName': 'ANA', 'identification': '1001', 'totalDocuments': 3,
                     'subtotal': 350000, 'discount': 0, 'total': 350000}]

    resp = client.get('/api/facts-freeze/backup.xlsx?year=2026', headers=h())
    assert resp.status_code == 200
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ['Información', 'Facturas', 'Prendas vendidas', 'Prendas de anuladas']
    assert wb['Facturas'].max_row == 1 + 4 and wb['Prendas vendidas'].max_row == 1 + 3


def test_solo_admin(app, client):
    from app.services.jwt_service import JWTService
    with app.app_context():
        token = JWTService.generate_token(2, 'cajera@test.com', 'cajero', None)
    hh = {'Authorization': f'Bearer {token}'}
    assert client.get('/api/facts-freeze/status', headers=hh).status_code == 403
    assert client.post('/api/facts-freeze/freeze', headers=hh, json={}).status_code == 403
